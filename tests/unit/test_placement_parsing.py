"""placeinator.placement.parsing -- tabular row extraction from XLSX/PDF/
DOCX placement sheets. Fixtures are real files built with the same libraries
that write them (openpyxl, python-docx) rather than hand-crafted binary
blobs, so a parsing-format regression shows up here, not just at runtime.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from docx import Document

from placeinator.placement.parsing import (
    EmptyPlacementDocumentError,
    UnsupportedPlacementFormatError,
    parse_placement_sheet_bytes,
)


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _docx_table_bytes(rows: list[list[str]]) -> bytes:
    document = Document()
    table = document.add_table(rows=0, cols=len(rows[0]))
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_xlsx_header_row_and_data_rows_are_parsed():
    data = _xlsx_bytes(
        [
            ["Student Name", "Result", "Interview Date"],
            ["Jane Doe", "Selected for interview", "2026-08-25"],
            ["John Smith", "Not selected", "2026-08-25"],
        ]
    )
    rows = parse_placement_sheet_bytes(data, "xlsx")
    assert rows == [
        {
            "Student Name": "Jane Doe",
            "Result": "Selected for interview",
            "Interview Date": "2026-08-25",
        },
        {"Student Name": "John Smith", "Result": "Not selected", "Interview Date": "2026-08-25"},
    ]


def test_docx_table_is_parsed():
    data = _docx_table_bytes([["Candidate", "Status", "Date"], ["Alice", "Shortlisted", "25 Aug"]])
    rows = parse_placement_sheet_bytes(data, "docx")
    assert rows == [{"Candidate": "Alice", "Status": "Shortlisted", "Date": "25 Aug"}]


def test_a_blank_row_between_header_and_data_is_not_returned_as_a_row():
    data = _xlsx_bytes([["Candidate"], [None], ["Jane Doe"]])
    rows = parse_placement_sheet_bytes(data, "xlsx")
    assert rows == [{"Candidate": "Jane Doe"}]


def test_a_sheet_with_only_a_header_raises_empty_document_error():
    data = _xlsx_bytes([["Candidate", "Status"]])
    with pytest.raises(EmptyPlacementDocumentError):
        parse_placement_sheet_bytes(data, "xlsx")


def test_unsupported_format_is_rejected_with_a_clear_message():
    with pytest.raises(UnsupportedPlacementFormatError, match="unsupported placement sheet format"):
        parse_placement_sheet_bytes(b"...", "txt")  # type: ignore[arg-type]
