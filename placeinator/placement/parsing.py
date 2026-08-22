"""Extracts tabular rows from a placement sheet attachment (spec section 7).

Mirrors `placeinator.resumes.parsing`'s PDF/DOCX handling -- kept as a
separate copy rather than a shared import, matching
`placeinator.jobs.parsing`'s own precedent, since a placement sheet shares
nothing with a resume/JD beyond which underlying library reads which file
format.

Unlike resumes/jobs parsing, the useful unit here isn't a text string -- it's
rows of column-name -> cell-text pairs, since placement sheets are tabular
data that `placeinator.placement.headers` then normalizes into canonical
fields. Header text is returned exactly as it appears in the source; nothing
here decides what a header "means".
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import Literal

import openpyxl
import pdfplumber
from docx import Document

PlacementSourceFormat = Literal["xlsx", "pdf", "docx"]

SUPPORTED_PLACEMENT_FORMATS: tuple[PlacementSourceFormat, ...] = ("xlsx", "pdf", "docx")


class UnsupportedPlacementFormatError(ValueError):
    pass


class EmptyPlacementDocumentError(ValueError):
    """Raised when parsing succeeds but recovers no rows -- e.g. an empty
    sheet, or a table-less PDF/DOCX containing only prose."""


class OcrUnavailableError(ValueError):
    """A scanned/image attachment was detected, but OCR (Tesseract) isn't
    available in this environment. Deliberately deferred -- see
    placeinator/placement/__init__.py -- rather than silently skipped or
    crashing the sync."""


def parse_placement_sheet_bytes(
    data: bytes, source_format: PlacementSourceFormat
) -> list[dict[str, str]]:
    if source_format == "xlsx":
        rows = _parse_xlsx(data)
    elif source_format == "pdf":
        rows = _parse_pdf(data)
    elif source_format == "docx":
        rows = _parse_docx(data)
    else:
        raise UnsupportedPlacementFormatError(
            f"unsupported placement sheet format: {source_format!r} "
            f"(expected one of {SUPPORTED_PLACEMENT_FORMATS})"
        )

    if not rows:
        raise EmptyPlacementDocumentError(
            "no tabular rows found -- if this is a scanned/image document, "
            "OCR support is not available yet"
        )
    return rows


def _rows_from_table(table: Sequence[Sequence[object]]) -> list[dict[str, str]]:
    """Shared by the XLSX/PDF/DOCX paths: the first non-empty row is the
    header, every row after becomes one dict keyed by that header. A row
    that's entirely blank cells is dropped, not returned as an empty dict."""
    rows_iter = iter(table)
    header: list[str] | None = None
    for raw in rows_iter:
        candidate = [str(cell).strip() if cell is not None else "" for cell in raw]
        if any(candidate):
            header = candidate
            break
    if header is None:
        return []

    rows: list[dict[str, str]] = []
    for raw_row in rows_iter:
        cells = [str(cell).strip() if cell is not None else "" for cell in raw_row]
        row = {header[i]: cells[i] for i in range(min(len(header), len(cells)))}
        if any(row.values()):
            rows.append(row)
    return rows


def _parse_xlsx(data: bytes) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    return _rows_from_table(list(sheet.iter_rows(values_only=True)))


def _parse_pdf(data: bytes) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend(_rows_from_table(table))
    return rows


def _parse_docx(data: bytes) -> list[dict[str, str]]:
    document = Document(io.BytesIO(data))
    rows: list[dict[str, str]] = []
    for table in document.tables:
        grid = [[cell.text for cell in row.cells] for row in table.rows]
        rows.extend(_rows_from_table(grid))
    return rows
